"""Render an anchors file as a readable rubric document (.docx).

Replaces the lost generator behind anchors_v3_rubric.docx (2026-09-01b; an npm `docx`
scratch script that never reached the repo). Same convention: every field is compared with
the FROZEN v1 anchors.json at render time -- not from a hand-written list -- and a changed
field is marked, with its v1 wording shown beneath. On top of that, a field touched by a
v3.1 change (make_anchors_v3.V3_1_CHANGES) carries a tag naming the change and the ruling
behind it, and a DERIVED change says plainly that it is an inference awaiting Julie.

    python make_anchors_docx.py                         # v3.1 -> ../../anchors_v3_1_rubric.docx
    python make_anchors_docx.py --anchors X.json --out Y.docx

Needs python-docx (pinned in requirements.txt).
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import make_anchors_v3 as mav  # noqa: E402

DEFAULT_OUT = HERE.parent.parent / "anchors_v3_1_rubric.docx"
GREY = RGBColor(0x6B, 0x72, 0x80)
FLAG = RGBColor(0x8A, 0x5A, 0x00)
V31 = RGBColor(0x1F, 0x4E, 0x79)


def change_tags() -> dict[tuple, str]:
    """path -> tag text, for every v3.1 change."""
    tags = {}
    for ch in mav.V3_1_CHANGES:
        who = (f"derived from ruling {ch['ruling']}: MY INFERENCE, awaiting Julie's confirmation"
               if ch["derived"] else f"Julie's ruling {ch['ruling']}")
        tags[tuple(ch["path"])] = f"v3.1 change {ch['id']} ({who})"
    return tags


class Renderer:
    def __init__(self, doc_anchors: dict, v1: dict):
        self.a, self.v1 = doc_anchors, v1
        # v3.1 change tags describe THAT file; on any other anchors file they would be false.
        self.tags = change_tags() if doc_anchors.get("rubric_version") == mav.V3_1_VERSION else {}
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name, style.font.size = "Calibri", Pt(10.5)
        self.marked: list[tuple] = []   # paths that received a v3.1 tag, for the test

    # -- primitives ---------------------------------------------------------------
    def _run(self, para, text, *, bold=False, italic=False, color=None, size=None):
        r = para.add_run(text)
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        if size is not None:
            r.font.size = Pt(size)
        return r

    def _markers(self, para, path: tuple, now, was) -> None:
        """Append the change markers for one field to `para`, and the v1 wording below."""
        tag = self.tags.get(path)
        if tag:
            self._run(para, f"  [{tag}]", bold=True, color=FLAG if "INFERENCE" in tag else V31,
                      size=9)
            self.marked.append(path)
        if was is None and now is not None:
            self._run(para, "  [NEW since v1]", bold=True, color=V31, size=9)
        elif was is not None and was != now:
            self._run(para, "  [CHANGED since v1]", bold=True, color=V31, size=9)
            p = self.doc.add_paragraph()
            self._run(p, "Was (v1): ", italic=True, color=GREY, size=9)
            self._run(p, str(was), italic=True, color=GREY, size=9)

    def field(self, label: str, now: str, was, path: tuple) -> None:
        p = self.doc.add_paragraph()
        self._run(p, f"{label}. ", bold=True)
        self._run(p, now)
        self._markers(p, path, now, was)

    # -- sections -----------------------------------------------------------------
    def header(self) -> None:
        self.doc.add_heading("Solverah company culture rubric", level=0)
        p = self.doc.add_paragraph()
        self._run(p, f"{self.a['rubric_version']}", bold=True)
        self._run(p, "  --  PROPOSED, NOT ADOPTED. Round 1 was scored against the frozen "
                     "v1-frozen-2026-08-15 file, which is unchanged and is never rescored.")
        p = self.doc.add_paragraph()
        self._run(p, "How to read the markers: ", bold=True)
        self._run(p, "a field whose wording differs from the frozen v1 rubric is marked "
                     "CHANGED, with the v1 wording shown beneath in grey; something v1 did not "
                     "have is marked NEW. A tag naming a v3.1 change (R1, R2a, ...) marks "
                     "where one of Julie's 2026-09-14 rulings landed; a tag that says MY "
                     "INFERENCE is an edit derived from a ruling that she has not yet "
                     "confirmed. ANCHORS_V3_1_CHANGELOG.md lists every change.")

    def gates(self) -> None:
        self.doc.add_heading("Before scoring -- the gates", level=1)
        for i, g in enumerate(self.a["gates"]):
            was = self.v1["gates"][i] if i < len(self.v1["gates"]) else None
            p = self.doc.add_paragraph(style="List Number")
            self._run(p, g)
            self._markers(p, ("gates", i), g, was)

    def confidence(self) -> None:
        self.doc.add_heading("Confidence bands (about the evidence, not about you)", level=1)
        v1c = {b["value"]: b for b in self.v1["confidence_bands"]}
        for b in self.a["confidence_bands"]:
            p = self.doc.add_paragraph()
            self._run(p, f"{b['value']} {b['label']}: ", bold=True)
            self._run(p, b["text"])
            self._markers(p, ("confidence_bands", b["value"], "text"), b["text"],
                          v1c.get(b["value"], {}).get("text"))

    def insufficient(self) -> None:
        ie = self.a.get("insufficient_evidence")
        if not ie:
            return
        self.doc.add_heading(f"{ie['label']} -- no score", level=1)
        p = self.doc.add_paragraph()
        self._run(p, ie["text"])
        self._markers(p, ("insufficient_evidence",), ie, self.v1.get("insufficient_evidence"))

    def agreement(self) -> None:
        self.doc.add_heading("How rater agreement is read", level=1)
        for k, v in self.a["agreement_rule"].items():
            p = self.doc.add_paragraph(style="List Bullet")
            self._run(p, f"{k.replace('_', ' ')}: ", bold=True)
            self._run(p, v)
            self._markers(p, ("agreement_rule", k), v, self.v1["agreement_rule"].get(k))

    def dimensions(self) -> None:
        v1d = {d["key"]: d for d in self.v1["dimensions"]}
        for d in self.a["dimensions"]:
            old = v1d[d["key"]]
            self.doc.add_page_break()
            self.doc.add_heading(f"{d['name']}:  {d['low_pole']} (0.1)  ->  {d['high_pole']} (0.9)",
                                 level=1)
            if d["low_pole"] != old["low_pole"] or d["high_pole"] != old["high_pole"]:
                p = self.doc.add_paragraph()
                self._run(p, f"Poles CHANGED since v1; v1 read: {old['low_pole']} -> "
                             f"{old['high_pole']}", italic=True, color=GREY, size=9)
            for label, key in (("Measures", "measures"), ("Where to look", "where_to_look"),
                               ("Keep distinct", "keep_distinct")):
                self.field(label, d[key], old[key], ("dimensions", d["key"], key))
            ob = {b["value"]: b for b in old["bands"]}
            table = self.doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, ("Band", "Reads as", "Description")):
                cell.text = text
            for b in d["bands"]:
                row = table.add_row().cells
                row[0].text = f"{b['value']}"
                row[1].text = b["reads_as"]
                para = row[2].paragraphs[0]
                self._run(para, b["text"])
                path = ("dimensions", d["key"], "bands", b["value"], "text")
                tag = self.tags.get(path)
                if tag:
                    self._run(para, f"  [{tag}]", bold=True, size=9,
                              color=FLAG if "INFERENCE" in tag else V31)
                    self.marked.append(path)
                was = ob[b["value"]]["text"]
                if was != b["text"]:
                    self._run(para, "  [CHANGED since v1]", bold=True, color=V31, size=9)
                    p2 = row[2].add_paragraph()
                    self._run(p2, "Was (v1): " + was, italic=True, color=GREY, size=9)
                if ob[b["value"]]["reads_as"] != b["reads_as"]:
                    self._run(row[1].paragraphs[0], f" (v1: {ob[b['value']]['reads_as']})",
                              italic=True, color=GREY, size=9)

    def render(self) -> Document:
        self.header()
        self.gates()
        self.confidence()
        self.insufficient()
        self.agreement()
        self.dimensions()
        return self.doc


def build(anchors_path: Path, out_path: Path) -> list[tuple]:
    """Write the docx; return the v3.1 change paths that were tagged in it."""
    v1 = json.loads(mav.V1_PATH.read_text(encoding="utf-8"))
    doc_anchors = json.loads(Path(anchors_path).read_text(encoding="utf-8"))
    r = Renderer(doc_anchors, v1)
    r.render().save(str(out_path))
    return r.marked


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--anchors", type=Path, default=mav.V3_1_PATH)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    a = ap.parse_args(argv)
    marked = build(a.anchors, a.out)
    print(f"wrote {a.out.name}: {len(marked)} v3.1 change tags")
    return 0


if __name__ == "__main__":
    sys.exit(main())
