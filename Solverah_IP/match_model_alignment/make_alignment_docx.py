"""Bundle the match-model alignment response (00..06) into one .docx for Julie.

House pattern: rater_app/make_anchors_docx.py (python-docx, Calibri 10.5). The markdown
files in this folder are the source of truth; this renders a small, known subset of
markdown (headings, paragraphs, bullets, numbered items, pipe tables, one fenced diagram,
bold / italic / code spans) into plain Word formatting, so no literal markdown symbol
reaches the document. Anything outside that subset is rendered as plain text.

    python make_alignment_docx.py                 # -> Match_Model_Alignment_Response_2026-10.docx
    python make_alignment_docx.py --out X.docx

Needs python-docx (pinned in requirements.txt).
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
DEFAULT_OUT = HERE / "Match_Model_Alignment_Response_2026-10.docx"
SOURCES = [
    "00_RESPONSE_MEMO.md",
    "01_ARCHITECTURE_MAP.md",
    "02_OUTCOME_DEFINITIONS.md",
    "03_SIX_TRAIT_BOUNDARY.md",
    "04_VALIDATION_ROADMAP.md",
    "05_USER_JOURNEYS.md",
    "06_COMPOUNDING_PAIRS_DRAFT.md",
]
GREY = RGBColor(0x6B, 0x72, 0x80)

# bold, then italic, then code. Order matters: ** must win over *.
INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`")
BULLET = re.compile(r"^(\s*)[-*] (.*)$")
NUMBERED = re.compile(r"^(\s*)\d+\. (.*)$")
TABLE_SEP = re.compile(r"^\|?\s*:?-{3,}")


def add_inline(para, text: str, *, size: float | None = None) -> None:
    """Render bold / italic / code spans as runs; markers never reach the document."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            if size:
                r.font.size = Pt(size)
        bold, italic, code = m.groups()
        r = para.add_run(bold or italic or code)
        r.bold = bold is not None
        r.italic = italic is not None
        if code is not None:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        if size:
            r.font.size = Pt(size)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    i, n = 0, len(lines)
    para_buf: list[str] = []

    def flush() -> None:
        if para_buf:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # closing fence
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name, r.font.size = "Consolas", Pt(7.5)
            continue
        if not stripped:
            flush()
            i += 1
            continue
        if stripped == "---":
            flush()
            i += 1
            continue
        h = re.match(r"^(#{1,6}) (.*)$", stripped)
        if h:
            flush()
            level = len(h.group(1))
            heading = doc.add_heading(level=0 if level == 1 else min(level - 1, 3))
            add_inline(heading, h.group(2))
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not TABLE_SEP.match(lines[i].strip().lstrip("|").strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncols):
                    cell_text = row[ci] if ci < len(row) else ""
                    para = cells[ci].paragraphs[0]
                    add_inline(para, f"**{cell_text}**" if ri == 0 and cell_text else cell_text,
                               size=9.5)
            doc.add_paragraph()
            continue
        b = BULLET.match(line)
        nm = NUMBERED.match(line)
        if b or nm:
            flush()
            m = b or nm
            item = [m.group(2)]
            i += 1
            # continuation lines: indented, non-blank, not a new list item
            while (i < n and lines[i].strip() and lines[i].startswith("  ")
                   and not BULLET.match(lines[i]) and not NUMBERED.match(lines[i])):
                item.append(lines[i].strip())
                i += 1
            if b:
                p = doc.add_paragraph(style="List Bullet 2" if m.group(1) else "List Bullet")
                add_inline(p, " ".join(item))
            else:
                # Word's "List Number" keeps counting across the whole document, so every
                # list after the first would start at 4, 9, ... Write the number instead.
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-14)
                number = line.strip().split(".", 1)[0]
                add_inline(p, f"{number}. " + " ".join(item))
            continue
        para_buf.append(line)
        i += 1
    flush()


def build(out_path: Path, sources: list[str] = SOURCES) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)

    doc.add_heading("Solverah: Product and Match Model Alignment", level=0)
    p = doc.add_paragraph()
    r = p.add_run("Response to Julie's October 2026 brief. Working draft, 2026-10-13. "
                  "Seven short documents; each stands alone, so read in any order and stop "
                  "whenever you like. Everything here is yours to edit.")
    r.italic = True
    r.font.color.rgb = GREY
    for idx, name in enumerate(sources):
        doc.add_page_break()
        render_markdown(doc, (HERE / name).read_text(encoding="utf-8"))
    doc.save(str(out_path))
    return doc


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args(argv)
    build(args.out)
    print(f"wrote {args.out.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
